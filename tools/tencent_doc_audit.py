#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
標準化部門 - 專案進度週報自動稽核工具
功能：
1. 讀取匯出的騰訊在線試算表 Excel 檔案。
2. 自動稽核顧問維護狀況：
   - 檢查哪些專案的「本週進度」或「執行狀況」是空白的（未維護週報）。
   - 檢查「預估上線/結案時間」是否逾期但未更新。
   - 彙整各顧問的專案達成與進度。
3. 產生稽核報告 Markdown 檔案並存入 Obsidian 的 100_Standardization/110_Audits/ 目錄下。
"""

import os
import sys
import re
from datetime import datetime

# ==================== 預設設定區 ====================
OBSIDIAN_VAULT_PATH = r"d:\AI\SecondBrain"
# ===================================================

# Ensure UTF-8 output
if sys.version_info >= (3, 7):
    sys.stdout.reconfigure(encoding='utf-8')

def find_column_index(headers, keywords):
    """根據關鍵字自動尋找最匹配的欄位索引"""
    for keyword in keywords:
        for idx, header in enumerate(headers):
            if header and keyword.lower() in str(header).lower():
                return idx
    return None

def analyze_excel(file_path):
    try:
        import openpyxl
    except ImportError:
        print("[錯誤] 請先安裝 openpyxl: pip install openpyxl")
        return None
        
    if not os.path.exists(file_path):
        print(f"[錯誤] 找不到檔案: {file_path}")
        return None
        
    wb = openpyxl.load_workbook(file_path, data_only=True)
    # 預設使用第一個工作表，或者名稱包含 "預估" 或 "2026" 的工作表
    sheet = wb.active
    for name in wb.sheetnames:
        if "預估" in name or "2026" in name or "週報" in name:
            sheet = wb[name]
            break
            
    print(f"[系統] 正在讀取工作表: {sheet.title} ...")
    
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("[錯誤] 試算表內容為空")
        return None
        
    # 尋找表頭行 (通常是包含 "客戶" 或 "顧問" 的前幾行)
    header_row_idx = 0
    for idx, row in enumerate(rows[:5]):
        row_str = " ".join([str(val) for val in row if val is not None])
        if "客戶" in row_str or "顧問" in row_str or "專案" in row_str:
            header_row_idx = idx
            break
            
    headers = [str(val).strip() if val is not None else "" for val in rows[header_row_idx]]
    data_rows = rows[header_row_idx + 1:]
    
    # 自動定位關鍵欄位
    col_client = find_column_index(headers, ["客戶名稱", "客戶", "專案名稱", "專案"])
    col_consultant = find_column_index(headers, ["負責顧問", "顧問", "負責人", "PM"])
    col_stage = find_column_index(headers, ["當前階段", "專案階段", "階段", "狀態", "當前狀態"])
    col_progress = find_column_index(headers, ["本週進度", "執行狀況", "進度", "本週工作", "狀況"])
    col_eta = find_column_index(headers, ["預估上線", "上線時間", "預計時間", "交期", "ETA"])
    col_warning = find_column_index(headers, ["異常說明", "提醒", "備註", "異常"])
    
    # 輸出欄位匹配狀況
    print("[系統] 欄位定位結果:")
    print(f"  - 客戶名稱: {'第 ' + str(col_client+1) + ' 欄' if col_client is not None else '未找到 (預設第1欄)'}")
    print(f"  - 負責顧問: {'第 ' + str(col_consultant+1) + ' 欄' if col_consultant is not None else '未找到 (預設第2欄)'}")
    print(f"  - 當前階段: {'第 ' + str(col_stage+1) + ' 欄' if col_stage is not None else '未找到'}")
    print(f"  - 進度說明: {'第 ' + str(col_progress+1) + ' 欄' if col_progress is not None else '未找到'}")
    print(f"  - 預估時間: {'第 ' + str(col_eta+1) + ' 欄' if col_eta is not None else '未找到'}")
    print(f"  - 異常備註: {'第 ' + str(col_warning+1) + ' 欄' if col_warning is not None else '未找到'}")
    
    # 設定預設欄位
    col_client = col_client if col_client is not None else 0
    col_consultant = col_consultant if col_consultant is not None else 1
    
    report_data = {
        'total_projects': 0,
        'unmaintained_list': [], # 週報未維護名單
        'delayed_list': [],      # 逾期未上線/結案專案
        'warnings_list': [],     # 有異常警報的專案
        'consultant_summary': {}, # 顧問維護統計
        'all_projects': []
    }
    
    for r_idx, row in enumerate(data_rows):
        # 排除整列空白
        if not any(val is not None for val in row):
            continue
            
        client = str(row[col_client]).strip() if row[col_client] is not None else ""
        # 排除表頭重複或無效行
        if not client or client == "None" or "客戶" in client:
            continue
            
        consultant = str(row[col_consultant]).strip() if row[col_consultant] is not None else "未指定"
        stage = str(row[col_stage]).strip() if (col_stage is not None and row[col_stage] is not None) else "未知"
        progress = str(row[col_progress]).strip() if (col_progress is not None and row[col_progress] is not None) else ""
        eta = str(row[col_eta]).strip() if (col_eta is not None and row[col_eta] is not None) else ""
        warning = str(row[col_warning]).strip() if (col_warning is not None and row[col_warning] is not None) else ""
        
        report_data['total_projects'] += 1
        
        # 統計各顧問專案數量與維護狀況
        if consultant not in report_data['consultant_summary']:
            report_data['consultant_summary'][consultant] = {'total': 0, 'maintained': 0, 'unmaintained': 0}
        report_data['consultant_summary'][consultant]['total'] += 1
        
        is_maintained = True
        # 稽核 1：進度說明是否空白
        if not progress or progress.upper() in ["NONE", "無", "未更新"]:
            report_data['unmaintained_list'].append({
                'client': client,
                'consultant': consultant,
                'stage': stage
            })
            report_data['consultant_summary'][consultant]['unmaintained'] += 1
            is_maintained = False
        else:
            report_data['consultant_summary'][consultant]['maintained'] += 1
            
        # 稽核 2：預估時間逾期判斷
        is_delayed = False
        if eta and eta != "None":
            try:
                # 嘗試解析常見的日期格式 (YYYY-MM-DD 或 YYYY/MM/DD)
                date_str = eta.split(" ")[0] # 排除時間部分
                date_str = re.sub(r'[^0-9/-]', '', date_str) # 移除非數字
                
                # 嘗試多種日期格式
                eta_date = None
                for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y-%m', '%Y/%m'):
                    try:
                        eta_date = datetime.strptime(date_str, fmt)
                        break
                    except ValueError:
                        continue
                
                if eta_date and eta_date < datetime.now() and "結案" not in stage and "上線" not in stage:
                    report_data['delayed_list'].append({
                        'client': client,
                        'consultant': consultant,
                        'stage': stage,
                        'eta': eta
                    })
                    is_delayed = True
            except Exception:
                pass # 無法解析日期則跳過
                
        # 稽核 3：異常警報
        if warning and warning.upper() not in ["NONE", "無", "正常", "無異常"]:
            report_data['warnings_list'].append({
                'client': client,
                'consultant': consultant,
                'stage': stage,
                'warning': warning
            })
            
        report_data['all_projects'].append({
            'client': client,
            'consultant': consultant,
            'stage': stage,
            'progress': progress if progress else "*未填寫*",
            'eta': eta if eta else "未填寫",
            'warning': warning if warning else "正常",
            'is_maintained': is_maintained,
            'is_delayed': is_delayed
        })
        
    return report_data

def generate_report(report_data, vault_path):
    today_str = datetime.now().strftime('%Y-%m-%d')
    today_display = datetime.now().strftime('%Y年%m月%d日')
    
    report_dir = os.path.join(vault_path, "100_Standardization", "110_Audits")
    os.makedirs(report_dir, exist_ok=True)
    
    filepath = os.path.join(report_dir, f"專案進度稽核報告_{today_str}.md")
    
    content = []
    content.append(f"# 📊 專案進度維護稽核報告 ({today_display})\n")
    content.append(f"- **稽核時間**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    content.append(f"- **標籤**: #audit/progress #report/weekly #standardization\n")
    content.append("---\n")
    
    # 1. 稽核摘要
    content.append("## 📌 稽核結果摘要")
    unmaintained_count = len(report_data['unmaintained_list'])
    delayed_count = len(report_data['delayed_list'])
    warning_count = len(report_data['warnings_list'])
    
    content.append(f"- **專案總數**: {report_data['total_projects']} 個專案")
    
    # 顯示高亮警報
    if unmaintained_count > 0:
        content.append(f"- ⚠️ **週報未更新專案**: =={unmaintained_count}== 個 (請督促顧問更新)")
    else:
        content.append("- ✅ **週報維護率**: 100% (所有專案皆已填寫進度)")
        
    if delayed_count > 0:
        content.append(f"- 🚨 **時間逾期未上線/結案專案**: =={delayed_count}== 個 (請確認專案是否延期)")
        
    if warning_count > 0:
        content.append(f"- 💥 **異常通報專案**: =={warning_count}== 個 (需主管介入協調)")
    content.append("")
    
    # 2. 顧問維護率統計
    content.append("### 👤 顧問週報維護率統計")
    content.append("| 顧問名稱 | 負責專案數 | 已維護 | 未維護 | 維護達成率 |")
    content.append("| :--- | :--- | :--- | :--- | :--- |")
    for name, stat in report_data['consultant_summary'].items():
        rate = (stat['maintained'] / stat['total']) * 100
        rate_str = f"{rate:.1f}%"
        # 低於 100% 標紅標記
        if rate < 100:
            rate_str = f"=={rate_str}=="
        content.append(f"| {name} | {stat['total']} | {stat['maintained']} | {stat['unmaintained']} | {rate_str} |")
    content.append("")
    
    # 3. 未維護週報詳細清單
    if report_data['unmaintained_list']:
        content.append("## ⚠️ 週報未更新專案清單")
        content.append("| 專案/客戶名稱 | 負責顧問 | 當前階段 |")
        content.append("| :--- | :--- | :--- |")
        for item in report_data['unmaintained_list']:
            content.append(f"| [[{item['client']}_基本資訊]] | [[{item['consultant']}]] | {item['stage']} |")
        content.append("")
        
    # 4. 逾期未更新清單
    if report_data['delayed_list']:
        content.append("## 🚨 預估日期逾期警報")
        content.append("> [!WARNING]")
        content.append("> 以下專案已超過預計上線/結案時間，但專案狀態仍為進行中。請確認實際進度並要求顧問修正預計交期。")
        content.append("")
        content.append("| 專案/客戶名稱 | 負責顧問 | 當前階段 | 原預估日期 |")
        content.append("| :--- | :--- | :--- | :--- |")
        for item in report_data['delayed_list']:
            content.append(f"| [[{item['client']}_基本資訊]] | [[{item['consultant']}]] | {item['stage']} | =={item['eta']}== |")
        content.append("")
        
    # 5. 異常說明清單
    if report_data['warnings_list']:
        content.append("## 💥 專案異常/主管關注警報")
        content.append("| 專案/客戶名稱 | 負責顧問 | 當前階段 | 異常說明與提醒 |")
        content.append("| :--- | :--- | :--- | :--- |")
        for item in report_data['warnings_list']:
            content.append(f"| [[{item['client']}_基本資訊]] | [[{item['consultant']}]] | {item['stage']} | <mark>{item['warning']}</mark> |")
        content.append("")
        
    # 6. 所有專案詳細清單
    content.append("## 📋 所有專案稽核細節")
    content.append("| 專案/客戶名稱 | 顧問 | 階段 | 週報更新 | 預計交期 | 異常/備註 |")
    content.append("| :--- | :--- | :--- | :---: | :--- | :--- |")
    for item in report_data['all_projects']:
        m_status = "✅" if item['is_maintained'] else "❌"
        eta_status = f"=={item['eta']}==" if item['is_delayed'] else item['eta']
        content.append(f"| [[{item['client']}_基本資訊]] | {item['consultant']} | {item['stage']} | {m_status} | {eta_status} | {item['warning']} |")
        
    content_text = "\n".join(content)
    
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content_text)
        
    print(f"[成功] 稽核週報已產出至: {filepath}")
    return filepath

def main():
    print("=== 專案週報進度稽核工具 ===")
    if len(sys.argv) < 2:
        print("[錯誤] 請提供騰訊文檔匯出的 Excel 檔案路徑！")
        print("用法: python tencent_doc_audit.py <Excel路徑>")
        sys.exit(1)
        
    excel_path = sys.argv[1]
    
    # 讀取分析
    report_data = analyze_excel(excel_path)
    if not report_data:
        print("[錯誤] 無法完成資料解析。")
        sys.exit(1)
        
    # 檢查 Obsidian 儲存路徑
    vault_path = OBSIDIAN_VAULT_PATH
    if not os.path.exists(vault_path):
        print(f"[提示] 找不到設定的 Vault 路徑: {vault_path}")
        print("將在當前目錄建立 'SecondBrain_Local' 資料夾進行測試歸檔。")
        vault_path = os.path.abspath("SecondBrain_Local")
        
    # 生成報告
    generate_report(report_data, vault_path)
    print("=== 稽核作業順利結束 ===")

if __name__ == "__main__":
    main()
